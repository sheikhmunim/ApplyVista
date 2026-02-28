# app/api.py

from __future__ import annotations

import io
import json
import traceback
from contextlib import asynccontextmanager
from pathlib import Path
from typing import Any, Dict, Optional

from fastapi import FastAPI, HTTPException
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import FileResponse, StreamingResponse
from pydantic import BaseModel, Field

from src.langgraph_pipeline import run_pipeline
from src.rag_pipeline import simple_chat_api
from src.job_scout_agent import run_scout, SCOUT_EXCEL
from src.scheduler import start_scheduler, scheduler

try:
    from src.config import OUT_DIR
except ImportError:
    OUT_DIR = Path("data/outputs")
    OUT_DIR.mkdir(parents=True, exist_ok=True)


# -----------------------------------------------------------
# Lifespan (startup / shutdown)
# -----------------------------------------------------------

@asynccontextmanager
async def lifespan(app: FastAPI):
    start_scheduler()
    yield
    scheduler.shutdown()


# -----------------------------------------------------------
# FastAPI app setup
# -----------------------------------------------------------

app = FastAPI(
    title="Job Assistant RAG API",
    description="RAG-based backend for job applications with LangGraph orchestration.",
    version="0.2.0",
    lifespan=lifespan,
)

app.add_middleware(
    CORSMiddleware,
    allow_origins=["http://localhost:3000", "http://127.0.0.1:3000"],
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)


# -----------------------------------------------------------
# Pydantic models
# -----------------------------------------------------------

class ChatRequest(BaseModel):
    message: str
    top_k: Optional[int] = None


class JDOptions(BaseModel):
    generate_skills: bool = True
    generate_cover: bool = True
    generate_emails: bool = True
    generate_ats: bool = True
    generate_top_choice: bool = True
    generate_short_recruiter_email: bool = True
    top_k: Optional[int] = Field(default=None, ge=1, le=20)


class JDRequest(BaseModel):
    jd_text: str = Field(..., min_length=30, max_length=30000)
    options: JDOptions = Field(default_factory=JDOptions)


class ExportRequest(BaseModel):
    text: str


# -----------------------------------------------------------
# Endpoints
# -----------------------------------------------------------

@app.get("/")
def root() -> Dict[str, str]:
    return {"status": "ok", "message": "Job Assistant RAG API v0.2 is running"}


@app.post("/jd")
def process_job_description(req: JDRequest) -> Dict[str, Any]:
    """
    Run the full LangGraph pipeline:
    - extract metadata (company, role)
    - generate documents (RAG pipeline)
    - search GitHub engineers
    - save Excel
    Returns all doc outputs + company_name, role_name, engineers, excel_filename.
    """
    try:
        state = run_pipeline(req.jd_text, options=req.options.model_dump())
        doc_result = state.get("doc_result", {})
        return {
            **doc_result,
            "company_name": state.get("company_name", ""),
            "role_name": state.get("role_name", ""),
            "engineers": state.get("engineers", []),
            "excel_filename": state.get("excel_filename", ""),
        }
    except Exception as e:
        traceback.print_exc()
        raise HTTPException(status_code=500, detail=str(e))


@app.post("/export/docx")
def export_docx(req: ExportRequest):
    """Convert text to a .docx file and return it for download."""
    try:
        from docx import Document

        doc = Document()
        for line in req.text.split("\n"):
            doc.add_paragraph(line)

        buf = io.BytesIO()
        doc.save(buf)
        buf.seek(0)

        return StreamingResponse(
            buf,
            media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            headers={"Content-Disposition": "attachment; filename=document.docx"},
        )
    except Exception as e:
        raise HTTPException(status_code=500, detail=str(e))


@app.post("/export/pdf")
def export_pdf(req: ExportRequest):
    """Convert text to a .pdf file and return it for download."""
    try:
        from fpdf import FPDF

        pdf = FPDF()
        pdf.add_page()
        pdf.set_font("Arial", size=11)
        for line in req.text.split("\n"):
            safe_line = line.encode("latin-1", errors="replace").decode("latin-1")
            pdf.multi_cell(0, 8, safe_line)

        buf = io.BytesIO(pdf.output(dest="S").encode("latin-1"))
        buf.seek(0)

        return StreamingResponse(
            buf,
            media_type="application/pdf",
            headers={"Content-Disposition": "attachment; filename=document.pdf"},
        )
    except Exception as e:
        raise HTTPException(status_code=500, detail=str(e))


@app.get("/download/{filename}")
def download_file(filename: str):
    """Serve a file from data/outputs/ by filename."""
    # Basic path traversal protection
    safe_name = Path(filename).name
    filepath = OUT_DIR / safe_name
    if not filepath.exists():
        raise HTTPException(status_code=404, detail=f"File '{safe_name}' not found")
    return FileResponse(str(filepath), filename=safe_name)


@app.get("/scout/jobs")
def get_scout_jobs() -> Dict[str, Any]:
    """Return all rows from job_scout.xlsx as JSON."""
    if not SCOUT_EXCEL.exists():
        return {"jobs": [], "total": 0}
    try:
        import openpyxl
        wb = openpyxl.load_workbook(str(SCOUT_EXCEL))
        ws = wb.active
        headers = [cell.value for cell in ws[1]]
        jobs = []
        for row in ws.iter_rows(min_row=2, values_only=True):
            jobs.append(dict(zip(headers, row)))
        return {"jobs": jobs, "total": len(jobs)}
    except Exception as e:
        raise HTTPException(status_code=500, detail=str(e))


@app.get("/scout/status")
def get_scout_status() -> Dict[str, Any]:
    """Return scheduler status: last run, next run, total jobs."""
    total = 0
    if SCOUT_EXCEL.exists():
        try:
            import openpyxl
            wb = openpyxl.load_workbook(str(SCOUT_EXCEL))
            ws = wb.active
            total = max(0, ws.max_row - 1)
        except Exception:
            pass

    job = scheduler.get_job("job_scout") if scheduler.running else None
    next_run = job.next_run_time.isoformat() if job and job.next_run_time else None

    return {
        "last_run": None,
        "next_run": next_run,
        "total_jobs": total,
        "scheduler_running": scheduler.running,
    }


@app.post("/scout/run")
def manual_scout_run() -> Dict[str, Any]:
    """Manually trigger the job scout agent."""
    try:
        result = run_scout()
        return result
    except Exception as e:
        traceback.print_exc()
        raise HTTPException(status_code=500, detail=str(e))


@app.post("/chat")
def chat(req: ChatRequest) -> Dict[str, Any]:
    k = req.top_k if req.top_k is not None else 6
    answer = simple_chat_api(req.message, k=k)
    return {"answer": answer}


@app.get("/health")
def health_check():
    return {"status": "healthy"}
