"""
job_scout_agent.py
Searches LinkedIn/Seek for relevant job postings, deduplicates,
appends to Excel, and sends email alerts.
"""

from __future__ import annotations

import json
import os
import smtplib
from datetime import datetime
from email.mime.text import MIMEText
from pathlib import Path
from typing import Any

from dotenv import load_dotenv

load_dotenv()

try:
    from .config import OUT_DIR, RAG_DIR
except ImportError:
    OUT_DIR = Path("data/outputs")
    RAG_DIR = Path("data/job_rag")

OUT_DIR.mkdir(parents=True, exist_ok=True)
RAG_DIR.mkdir(parents=True, exist_ok=True)

SCOUT_EXCEL = OUT_DIR / "job_scout.xlsx"
SEEN_JOBS_PATH = RAG_DIR / "seen_jobs.json"

SMTP_HOST = os.getenv("SMTP_HOST", "smtp.gmail.com")
SMTP_PORT = int(os.getenv("SMTP_PORT", "587"))
SMTP_USER = os.getenv("SMTP_USER", "")
SMTP_PASSWORD = os.getenv("SMTP_PASSWORD", "")
ALERT_EMAIL = os.getenv("ALERT_EMAIL", SMTP_USER)

SCOUT_ROLE = os.getenv("SCOUT_ROLE", "ML Engineer")
SCOUT_LOCATION = os.getenv("SCOUT_LOCATION", "Melbourne")


# ---------------------------------------------------------------------------
# Search query builder
# ---------------------------------------------------------------------------

def build_search_queries(scout_role: str = SCOUT_ROLE, location: str = SCOUT_LOCATION) -> list[str]:
    """Build 3-5 search queries for the job scout."""
    return [
        f"{scout_role} jobs {location} site:seek.com.au",
        f"{scout_role} jobs {location} site:au.linkedin.com/jobs",
        f"AI Engineer jobs {location} site:seek.com.au",
        f"Machine Learning Engineer jobs {location} site:seek.com.au",
        f"{scout_role} {location} hiring 2024",
    ]


# ---------------------------------------------------------------------------
# Search
# ---------------------------------------------------------------------------

def search_jobs(query: str) -> list[dict]:
    """
    Search for jobs using DuckDuckGo text search.
    Returns a list of {title, url, snippet, source} dicts.
    """
    try:
        from duckduckgo_search import DDGS
    except ImportError:
        return []

    results = []
    try:
        with DDGS() as ddgs:
            for r in ddgs.text(query, max_results=10):
                source = "seek" if "seek" in r.get("href", "") else "linkedin"
                results.append({
                    "title": r.get("title", ""),
                    "url": r.get("href", ""),
                    "snippet": r.get("body", ""),
                    "source": source,
                })
    except Exception:
        pass
    return results


# ---------------------------------------------------------------------------
# Scoring
# ---------------------------------------------------------------------------

def score_job_fit(snippet: str, user_skills: list[str]) -> int:
    """
    Keyword overlap between job snippet and user skills.
    Returns 0-100 score.
    """
    if not snippet or not user_skills:
        return 0
    snippet_lower = snippet.lower()
    matches = sum(1 for skill in user_skills if skill.lower() in snippet_lower)
    return min(100, int((matches / max(len(user_skills), 1)) * 100 * 3))


# ---------------------------------------------------------------------------
# Dedup
# ---------------------------------------------------------------------------

def _load_seen(path: Path) -> set[str]:
    if path.exists():
        try:
            return set(json.loads(path.read_text(encoding="utf-8")))
        except Exception:
            return set()
    return set()


def _save_seen(seen: set[str], path: Path) -> None:
    path.write_text(json.dumps(sorted(seen), indent=2), encoding="utf-8")


def filter_new_jobs(found: list[dict], seen_path: Path = SEEN_JOBS_PATH) -> list[dict]:
    """
    Return only jobs with URLs not in the seen set.
    Updates seen_jobs.json with new URLs.
    """
    seen = _load_seen(seen_path)
    new_jobs = [j for j in found if j.get("url") and j["url"] not in seen]
    seen.update(j["url"] for j in new_jobs)
    _save_seen(seen, seen_path)
    return new_jobs


# ---------------------------------------------------------------------------
# Excel persistence
# ---------------------------------------------------------------------------

def append_to_excel(new_jobs: list[dict], excel_path: Path = SCOUT_EXCEL) -> None:
    """
    Append new job rows to the persistent scout Excel file.
    Creates file with headers if it doesn't exist.
    """
    try:
        import openpyxl
    except ImportError:
        return

    headers = ["Title", "Company", "URL", "Fit Score", "Source", "Date Found", "Status"]

    if excel_path.exists():
        wb = openpyxl.load_workbook(str(excel_path))
        ws = wb.active
    else:
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Jobs"
        ws.append(headers)

    now_str = datetime.now().strftime("%Y-%m-%d %H:%M")
    for job in new_jobs:
        ws.append([
            job.get("title", ""),
            job.get("company", ""),
            job.get("url", ""),
            job.get("score", 0),
            job.get("source", ""),
            now_str,
            "New",
        ])

    wb.save(str(excel_path))


# ---------------------------------------------------------------------------
# Email alert
# ---------------------------------------------------------------------------

def send_email_alert(new_jobs: list[dict]) -> None:
    """Send a plain-text email listing new job findings via Gmail SMTP."""
    if not SMTP_USER or not SMTP_PASSWORD or not new_jobs:
        return

    lines = [f"Job Scout found {len(new_jobs)} new job(s):\n"]
    for i, job in enumerate(new_jobs, 1):
        lines.append(f"{i}. {job.get('title', 'Untitled')}")
        lines.append(f"   Score: {job.get('score', 0)}/100")
        lines.append(f"   Source: {job.get('source', '')}")
        lines.append(f"   URL: {job.get('url', '')}")
        lines.append("")

    body = "\n".join(lines)
    msg = MIMEText(body)
    msg["Subject"] = f"[Job Scout] {len(new_jobs)} new job(s) found"
    msg["From"] = SMTP_USER
    msg["To"] = ALERT_EMAIL

    try:
        with smtplib.SMTP(SMTP_HOST, SMTP_PORT) as server:
            server.ehlo()
            server.starttls()
            server.login(SMTP_USER, SMTP_PASSWORD)
            server.sendmail(SMTP_USER, [ALERT_EMAIL], msg.as_string())
    except Exception:
        pass


# ---------------------------------------------------------------------------
# Main orchestrator
# ---------------------------------------------------------------------------

def run_scout() -> dict[str, Any]:
    """
    Run the full job scout:
    1. Build search queries
    2. Search DuckDuckGo
    3. Score each result
    4. Filter for new jobs
    5. Append to Excel
    6. Send email alert

    Returns {new_jobs_count, total_jobs_count, timestamp}
    """
    try:
        from .profile_config import USER_PROFILE
        skills = USER_PROFILE.get("skills", [])
    except ImportError:
        skills = []

    queries = build_search_queries()
    all_found: list[dict] = []
    seen_urls: set[str] = set()

    for query in queries:
        results = search_jobs(query)
        for job in results:
            url = job.get("url", "")
            if url and url not in seen_urls:
                seen_urls.add(url)
                job["score"] = score_job_fit(job.get("snippet", ""), skills)
                all_found.append(job)

    new_jobs = filter_new_jobs(all_found)

    # Count total rows in Excel before appending
    total_before = 0
    if SCOUT_EXCEL.exists():
        try:
            import openpyxl
            wb = openpyxl.load_workbook(str(SCOUT_EXCEL))
            ws = wb.active
            total_before = max(0, ws.max_row - 1)  # subtract header row
        except Exception:
            pass

    if new_jobs:
        append_to_excel(new_jobs)
        send_email_alert(new_jobs)

    return {
        "new_jobs_count": len(new_jobs),
        "total_jobs_count": total_before + len(new_jobs),
        "timestamp": datetime.now().isoformat(),
    }
